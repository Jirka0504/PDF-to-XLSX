import re
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from pathlib import Path

import pdfplumber
from openpyxl import Workbook


APP_TITLE = "OMNIA PDF → XLSX převodník"


HEADERS = [
    "Interní kód zboží",
    "Název",
    "Množství",
    "Cena celkem",
    "Zkrácená poznámka",
    "Kód kombinované nomenklatury",
    "Země původu",
    "Hmotnost",
]


def normalize_text(text: str) -> str:
    text = text.replace("\u00a0", " ")
    text = text.replace("￾", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def clean_number(value: str) -> float:
    value = value.replace("€", "").replace("EUR", "").strip()
    value = value.replace(",", ".")
    return float(value)


def looks_like_code_only(line: str) -> bool:
    """
    Samostatný kód na řádku, např.:
    VEN149198350
    nebo
    SS-193757
    """
    line = normalize_text(line)
    if " PZ " in line:
        return False
    return bool(re.fullmatch(r"[A-Z0-9\-.]{4,}", line))


def parse_pdf(pdf_path: str):
    lines = []

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for line in text.splitlines():
                line = normalize_text(line)
                if line:
                    lines.append(line)

    items = []
    warnings = []

    # ---------------------------------------------------------
    # DETEKCE NOVÉHO TYPU OMNIA
    # ---------------------------------------------------------

    full_text = "\n".join(lines).lower()

    is_new_omnia = (
        "code cust. item description qty. u.m. price amount vat c." in full_text
        or (
            "shipment no." in full_text
            and "pcs" in full_text
            and "omnia components" in full_text
        )
    )

    # =========================================================
    # NOVÝ OMNIA FORMÁT
    # =========================================================

    if is_new_omnia:

        # běžný kompletní řádek:
        #
        # PLM002BO DRYER TUMBLE JOCKEY PULLEY-BOSCH 00632045
        # 5 Pcs 1,05 5,25 41

        new_row_re = re.compile(
            r"""
            ^
            (?P<code>[A-Z0-9.\-]+)
            \s+
            (?P<name>.+?)
            \s+
            (?P<qty>\d+)
            \s+Pcs
            \s+
            (?P<price>\d+(?:[.,]\d{2}))
            \s+
            (?P<total>\d+(?:[.,]\d{2}))
            \s+
            (?P<vat>\d+)
            $
            """,
            re.VERBOSE | re.IGNORECASE,
        )

        # samotný konec řádku:
        # 2 Pcs 1,35 2,70 41

        new_tail_re = re.compile(
            r"""
            ^
            (?P<qty>\d+)
            \s+Pcs
            \s+
            (?P<price>\d+(?:[.,]\d{2}))
            \s+
            (?P<total>\d+(?:[.,]\d{2}))
            \s+
            (?P<vat>\d+)
            $
            """,
            re.VERBOSE | re.IGNORECASE,
        )

        def decimal_new(value):
            return float(
                str(value)
                .replace(" ", "")
                .replace(".", "")
                .replace(",", ".")
            )

        skip_prefixes_new = (
            "Code Cust. Item",
            "Shipment No.",
            "Note:",
            "Company Stamp",
            "Parcel units:",
            "Gross weight",
            "Goods aspect:",
            "Vat C.",
            "Incoterm:",
            "Shipment due to:",
            "Shipping Agent:",
            "Subscr. No.:",
            "Truck. No.:",
            "Shipping Starting Date",
            "Driver signature:",
            "Addresse signature:",
            "SWIFT - Bank transfer",
            "No. ",
            "Via Travnik",
            "Company subject",
            "SDI Code:",
            "VAT Registration",
            "Ph:",
            "Omnia Components Srl",
            "KTS - AME",
            "Czech Republic",
            "Company:",
            "Karla Čapka",
            "Deliver to:",
            "Invoice",
            "Payment:",
            "Paym. Method:",
            "Net 60",
            "Bill-to Customer",
            "Due Dates:",
            "Bank:",
            "PAG",
            "TOTAL DOCUMENT",
            "Pursuant to",
            "Delivery At Place",
            "Fedex",
        )

        cleaned = []

        for raw_line in lines:
            line = normalize_text(raw_line)

            if not line:
                continue

            if any(
                line.startswith(prefix)
                for prefix in skip_prefixes_new
            ):
                continue

            cleaned.append(line)

        i = 0

        while i < len(cleaned):
            line = cleaned[i]

            # ---------------------------------------------
            # DOPRAVA - VYNECHAT
            # ---------------------------------------------

            if line.startswith("TRASP.EU.VEN"):
                i += 1
                continue

            if "Shipping Fees" in line:
                i += 1
                continue

            # ---------------------------------------------
            # 1) BĚŽNÝ KOMPLETNÍ ŘÁDEK
            # ---------------------------------------------

            m = new_row_re.match(line)

            if m:
                code = m.group("code").strip()

                if code == "TRASP.EU.VEN":
                    i += 1
                    continue

                items.append({
                    "code": code,
                    "name": m.group("name").strip(),
                    "qty": int(m.group("qty")),
                    "total": decimal_new(
                        m.group("total")
                    ),
                })

                i += 1
                continue

            # ---------------------------------------------
            # 2) VÍCEŘÁDKOVÁ POLOŽKA
            #
            # LFT003UN WASHING MACHINE RUBBER ...
            # SKL
            # 2 Pcs 1,35 2,70 41
            # ---------------------------------------------

            start = re.match(
                r"^(?P<code>[A-Z0-9.\-]+)\s+(?P<name>.+)$",
                line,
            )

            if start:
                code = start.group("code").strip()

                if code == "TRASP.EU.VEN":
                    i += 1
                    continue

                name_parts = [
                    start.group("name").strip()
                ]

                j = i + 1
                found = False

                while j < len(cleaned):
                    nxt = cleaned[j]

                    tail = new_tail_re.match(nxt)

                    if tail:
                        items.append({
                            "code": code,
                            "name": " ".join(
                                name_parts
                            ).strip(),
                            "qty": int(
                                tail.group("qty")
                            ),
                            "total": decimal_new(
                                tail.group("total")
                            ),
                        })

                        i = j + 1
                        found = True
                        break

                    # pokud narazíme na další zjevnou položku,
                    # ukončíme hledání
                    if new_row_re.match(nxt):
                        break

                    name_parts.append(nxt)
                    j += 1

                if found:
                    continue

            i += 1

        return items, warnings

    # =========================================================
    # STARÝ OMNIA FORMÁT
    # =========================================================

    pending_code = None

    full_row_re = re.compile(
        r"""
        ^
        (?P<code>[A-Z0-9\-.]+)
        \s+
        (?P<name>.+?)
        \s+
        (?P<qty>\d+)
        \s+PZ
        \s+
        (?P<price>\d+(?:[.,]\d{2}))
        \s+€
        \s+
        (?P<total>\d+(?:[.,]\d{2}))
        \s+€
        $
        """,
        re.VERBOSE | re.IGNORECASE,
    )

    cont_row_re = re.compile(
        r"""
        ^
        (?P<name>.+?)
        \s+
        (?P<qty>\d+)
        \s+PZ
        \s+
        (?P<price>\d+(?:[.,]\d{2}))
        \s+€
        \s+
        (?P<total>\d+(?:[.,]\d{2}))
        \s+€
        $
        """,
        re.VERBOSE | re.IGNORECASE,
    )

    skip_prefixes = (
        "Consegnare a:",
        "Spettabile:",
        "Vat code:",
        "TOTALE MERCE",
        "SCONTO %",
        "SPESE INCASSO",
        "TOTALE IMPONIBILE",
        "TOTALE IMPOSTA",
        "IMPONIBILE",
        "ALIQUOTA IVA",
        "TOTALE DOCUMENTO",
        "OMAGGI",
        "TOTALE DA PAGARE",
        "FATTURA",
        "OMNIA COMPONENTS",
        "Via Travnik",
        "Tel.",
        "C.F. E P.IVA",
        "Capitale Sociale",
        "N.Documento",
        "Data spedizione richiesta:",
        "http://www.omniacomponents.com",
        "PRODUCT CODE DESCRIPTION QUANTITY PREZZO SCONTO % IMPORTO TOTALE",
        "Spese di Trasporto UE Vendita - Shipping Fees",
    )

    for raw_line in lines:
        line = normalize_text(raw_line)

        # znak z PDF, který reprezentuje rozdělenou pomlčku
        line = line.replace("￾", "-")

        if any(
            line.startswith(prefix)
            for prefix in skip_prefixes
        ):
            continue

        if (
            line.startswith("KTS - AME")
            or line.startswith("Karla Čapka")
            or line.startswith("500 02")
        ):
            continue

        # ---------------------------------------------
        # 1) normální kompletní řádek
        # ---------------------------------------------

        m = full_row_re.match(line)

        if m:
            code = m.group("code").strip()

            items.append({
                "code": code,
                "name": m.group("name").strip(),
                "qty": int(m.group("qty")),
                "total": clean_number(
                    m.group("total")
                ),
            })

            pending_code = None
            continue

        # ---------------------------------------------
        # 2) samostatný kód
        # ---------------------------------------------

        if looks_like_code_only(line):
            pending_code = line
            continue

        # ---------------------------------------------
        # 3) pokračovací řádek bez kódu
        # ---------------------------------------------

        m2 = cont_row_re.match(line)

        if m2 and pending_code:
            items.append({
                "code": pending_code.strip(),
                "name": m2.group("name").strip(),
                "qty": int(m2.group("qty")),
                "total": clean_number(
                    m2.group("total")
                ),
            })

            pending_code = None
            continue

        if " PZ " in line or " Pcs " in line:
            warnings.append(
                f"Nepodařilo se naparsovat řádek: {line}"
            )

    return items, warnings


def save_xlsx(output_path: str, items: list[dict]):
    wb = Workbook()
    ws = wb.active
    ws.title = "List1"

    # hlavička
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    # data
    row = 2
    for item in items:
        ws.cell(row=row, column=1, value=item["code"])
        ws.cell(row=row, column=2, value=item["name"])
        ws.cell(row=row, column=3, value=item["qty"])
        ws.cell(row=row, column=4, value=item["total"])
        ws.cell(row=row, column=5, value=None)
        ws.cell(row=row, column=6, value=None)
        ws.cell(row=row, column=7, value=None)
        ws.cell(row=row, column=8, value=None)
        row += 1

    wb.save(output_path)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("900x560")

        self.pdf_path = tk.StringVar()
        self.output_path = tk.StringVar()

        self.preview_text = None

        self._build_ui()

    def _build_ui(self):
        frm = ttk.Frame(self, padding=12)
        frm.pack(fill="both", expand=True)

        ttk.Label(frm, text="Vstupní PDF:").grid(row=0, column=0, sticky="w", pady=4)
        ttk.Entry(frm, textvariable=self.pdf_path, width=90).grid(row=0, column=1, sticky="ew", pady=4)
        ttk.Button(frm, text="Vybrat…", command=self.pick_pdf).grid(row=0, column=2, padx=6)

        ttk.Label(frm, text="Výstupní XLSX:").grid(row=1, column=0, sticky="w", pady=4)
        ttk.Entry(frm, textvariable=self.output_path, width=90).grid(row=1, column=1, sticky="ew", pady=4)
        ttk.Button(frm, text="Uložit jako…", command=self.pick_output).grid(row=1, column=2, padx=6)

        ttk.Separator(frm, orient="horizontal").grid(row=2, column=0, columnspan=3, sticky="ew", pady=10)

        note = (
            "Poznámka: Pokud je kód zboží rozdělený do dvou řádků, aplikace vezme samostatný kód\n"
            "z jednoho řádku a spojí ho s následujícím řádkem s názvem, množstvím a cenou."
        )
        ttk.Label(frm, text=note, foreground="#555").grid(row=3, column=0, columnspan=3, sticky="w", pady=(0, 10))

        self.preview_text = tk.Text(frm, height=18, wrap="none")
        self.preview_text.grid(row=4, column=0, columnspan=3, sticky="nsew")

        btns = ttk.Frame(frm)
        btns.grid(row=5, column=0, columnspan=3, sticky="e", pady=12)
        ttk.Button(btns, text="Náhled", command=self.preview).pack(side="left", padx=4)
        ttk.Button(btns, text="Převést", command=self.convert).pack(side="left", padx=4)
        ttk.Button(btns, text="Konec", command=self.destroy).pack(side="left", padx=4)

        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(4, weight=1)

    def pick_pdf(self):
        path = filedialog.askopenfilename(filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")])
        if path:
            self.pdf_path.set(path)
            if not self.output_path.get().strip():
                self.output_path.set(str(Path(path).with_suffix(".xlsx")))

    def pick_output(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
        if path:
            self.output_path.set(path)

    def preview(self):
        pdf = self.pdf_path.get().strip()
        if not pdf:
            messagebox.showerror("Chyba", "Vyber vstupní PDF.")
            return

        try:
            items, warnings = parse_pdf(pdf)
        except Exception as e:
            messagebox.showerror("Chyba", str(e))
            return

        self.preview_text.delete("1.0", "end")
        self.preview_text.insert("1.0", f"Nalezeno položek: {len(items)}\n\n")

        for item in items[:30]:
            self.preview_text.insert(
                "end",
                f"{item['code']} | {item['name']} | qty={item['qty']} | total={item['total']}\n"
            )

        if warnings:
            self.preview_text.insert("end", "\n--- VAROVÁNÍ ---\n")
            for w in warnings[:20]:
                self.preview_text.insert("end", w + "\n")

    def convert(self):
        pdf = self.pdf_path.get().strip()
        out = self.output_path.get().strip()

        if not pdf:
            messagebox.showerror("Chyba", "Vyber vstupní PDF.")
            return
        if not out:
            messagebox.showerror("Chyba", "Vyber výstupní XLSX.")
            return

        try:
            items, warnings = parse_pdf(pdf)
            save_xlsx(out, items)
        except Exception as e:
            messagebox.showerror("Chyba", str(e))
            return

        msg = f"Hotovo.\nUloženo do:\n{out}\n\nPočet položek: {len(items)}"
        if warnings:
            msg += f"\n\nVarování: {len(warnings)}"
        messagebox.showinfo("Hotovo", msg)


if __name__ == "__main__":
    App().mainloop()
